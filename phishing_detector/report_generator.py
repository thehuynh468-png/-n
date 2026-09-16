"""
report_generator.py
--------------------
Tạo báo cáo Word (.docx) tổng kết kết quả phát hiện phishing.
"""

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logging.warning("python-docx không khả dụng. Không tạo được báo cáo Word.")

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────
# Styling helpers
# ─────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str):
    """Đặt màu nền ô bảng."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _set_cell_bold(cell, bold: bool = True):
    for para in cell.paragraphs:
        for run in para.runs:
            run.bold = bold


def _add_heading(doc, text: str, level: int = 1, color_hex: str = "1F3864"):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(
            int(color_hex[:2], 16),
            int(color_hex[2:4], 16),
            int(color_hex[4:], 16),
        )
    return heading


def _add_comparison_table(doc, df: pd.DataFrame):
    """Thêm bảng so sánh mô hình vào document."""
    cols = list(df.columns)
    table = doc.add_table(rows=1 + len(df), cols=len(cols))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr = table.rows[0]
    for i, col in enumerate(cols):
        cell = hdr.cells[i]
        cell.text = col
        _set_cell_bg(cell, "1F3864")
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)

    # Data rows
    best_f1 = df["F1-Score"].max() if "F1-Score" in df.columns else None
    for row_idx, row_data in df.iterrows():
        row = table.rows[row_idx + 1]
        is_best = ("F1-Score" in df.columns and row_data["F1-Score"] == best_f1)

        for col_idx, col in enumerate(cols):
            cell = row.cells[col_idx]
            val = row_data[col]
            cell.text = str(val) if not isinstance(val, float) else f"{val:.4f}"

            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(9)
                    if is_best:
                        run.bold = True

            if is_best:
                _set_cell_bg(cell, "D6E4F0")

        row_bg = "F2F2F2" if row_idx % 2 == 0 else "FFFFFF"
        if not is_best:
            for col_idx in range(len(cols)):
                _set_cell_bg(row.cells[col_idx], row_bg)


def _add_stats_table(doc, stats_df: pd.DataFrame):
    """Thêm bảng thống kê mô tả."""
    stats_df = stats_df.round(3)
    cols = ["Feature"] + list(stats_df.columns)
    table = doc.add_table(rows=1 + len(stats_df), cols=len(cols))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0]
    for i, col in enumerate(cols):
        cell = hdr.cells[i]
        cell.text = col
        _set_cell_bg(cell, "2E75B6")
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8.5)

    for row_idx, (feat_name, row_data) in enumerate(stats_df.iterrows()):
        row = table.rows[row_idx + 1]
        row_bg = "EBF3FB" if row_idx % 2 == 0 else "FFFFFF"

        row.cells[0].text = str(feat_name)
        _set_cell_bg(row.cells[0], row_bg)
        for para in row.cells[0].paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.size = Pt(8.5)

        for col_idx, col in enumerate(stats_df.columns, start=1):
            cell = row.cells[col_idx]
            cell.text = f"{row_data[col]:.3f}"
            _set_cell_bg(cell, row_bg)
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(8.5)


# ─────────────────────────────────────────────
# Main report generation
# ─────────────────────────────────────────────

def generate_report(
    output_path: str,
    results_df: pd.DataFrame,
    best_model: dict,
    stats_df: pd.DataFrame,
    total_urls: int,
    phishing_count: int,
    benign_count: int,
    feature_cols: list,
    dataset_path: str = "phishing_urls.csv",
):
    """
    Tạo file báo cáo Word.

    Args:
        output_path: Đường dẫn file .docx đầu ra.
        results_df: DataFrame kết quả so sánh mô hình.
        best_model: dict thông tin mô hình tốt nhất.
        stats_df: DataFrame thống kê mô tả đặc trưng lexical.
        total_urls: Tổng số URL trong dataset.
        phishing_count: Số URL phishing.
        benign_count: Số URL benign.
        feature_cols: Danh sách tên cột đặc trưng lexical.
        dataset_path: Tên file dataset.
    """
    if not DOCX_AVAILABLE:
        logger.error("python-docx chưa được cài đặt. Không thể tạo báo cáo.")
        return

    doc = Document()

    # ── Page margins ──
    from docx.oxml.ns import qn as _qn
    section = doc.sections[0]
    section.page_width  = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    # ── Default font ──
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # ══════════════════════════════════════════
    # Trang bìa
    # ══════════════════════════════════════════
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("BÁO CÁO PHÁT HIỆN URL LỪA ĐẢO (PHISHING)")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(31, 56, 100)
    run.font.name = "Times New Roman"

    doc.add_paragraph()

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_para.add_run("Hệ thống Machine Learning phân loại URL Phishing / Benign")
    sub_run.italic = True
    sub_run.font.size = Pt(13)
    sub_run.font.name = "Times New Roman"

    doc.add_paragraph()
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_para.add_run(f"Ngày tạo: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    date_run.font.size = Pt(11)
    date_run.font.name = "Times New Roman"

    doc.add_page_break()

    # ══════════════════════════════════════════
    # 1. Mô tả bài toán và dữ liệu
    # ══════════════════════════════════════════
    _add_heading(doc, "1. Mô tả bài toán và dữ liệu", level=1)

    doc.add_paragraph(
        "Bài toán phát hiện URL lừa đảo (phishing detection) là một bài toán phân loại nhị phân "
        "trong lĩnh vực an ninh mạng. Mục tiêu là xây dựng mô hình học máy có khả năng phân biệt "
        "URL độc hại (phishing) và URL an toàn (benign) dựa trên các đặc trưng được trích xuất "
        "từ chính chuỗi URL."
    )

    _add_heading(doc, "1.1. Thông tin dataset", level=2)

    info_table = doc.add_table(rows=4, cols=2)
    info_table.style = "Table Grid"
    items = [
        ("Tên file dataset", dataset_path),
        ("Tổng số URL", str(total_urls)),
        ("Số URL Phishing (label=1)", str(phishing_count)),
        ("Số URL Benign (label=0)", str(benign_count)),
    ]
    for i, (k, v) in enumerate(items):
        row = info_table.rows[i]
        row.cells[0].text = k
        row.cells[1].text = v
        _set_cell_bg(row.cells[0], "DEEAF1")
        for para in row.cells[0].paragraphs:
            for run in para.runs:
                run.bold = True

    doc.add_paragraph()

    # ══════════════════════════════════════════
    # 2. Đặc trưng đã trích xuất
    # ══════════════════════════════════════════
    _add_heading(doc, "2. Danh sách đặc trưng đã trích xuất", level=1)

    _add_heading(doc, "2.1. Nhóm 1 – Đặc trưng Lexical (chuỗi ký tự)", level=2)
    doc.add_paragraph(
        "Các đặc trưng này được trích xuất trực tiếp từ chuỗi URL mà không cần tra cứu mạng, "
        "do đó rất nhanh và phù hợp cho xử lý batch lớn."
    )

    lexical_features = [
        ("url_length",          "Tổng độ dài chuỗi URL"),
        ("domain_length",       "Độ dài phần domain + TLD"),
        ("num_dots",            "Số dấu chấm (.) trong URL"),
        ("num_slashes",         "Số dấu gạch chéo (/) trong URL"),
        ("num_question_marks",  "Số dấu hỏi chấm (?)"),
        ("num_equal_signs",     "Số dấu bằng (=)"),
        ("num_hyphens",         "Số dấu gạch ngang (-)"),
        ("num_underscores",     "Số dấu gạch dưới (_)"),
        ("num_at_signs",        "Số dấu @ – thường dùng để che giấu domain thật"),
        ("num_and_signs",       "Số dấu & trong query string"),
        ("path_depth",          "Số cấp đường dẫn (số / trong path)"),
        ("query_length",        "Độ dài chuỗi query (sau ?)"),
        ("num_query_params",    "Số tham số query"),
        ("has_sensitive_keyword","1 nếu có từ nhạy cảm (login, dangnahp, bank, ...)"),
        ("has_ip_address",      "1 nếu host là địa chỉ IPv4"),
        ("is_shortened",        "1 nếu dùng dịch vụ rút gọn URL (bit.ly, ...)"),
        ("uses_https",          "1 nếu dùng HTTPS"),
        ("num_subdomains",      "Số lượng subdomain (bỏ www)"),
    ]

    feat_table = doc.add_table(rows=1 + len(lexical_features), cols=2)
    feat_table.style = "Table Grid"
    feat_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, hdr_text in enumerate(["Tên đặc trưng", "Mô tả"]):
        cell = feat_table.rows[0].cells[i]
        cell.text = hdr_text
        _set_cell_bg(cell, "1F3864")
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

    for row_idx, (name, desc) in enumerate(lexical_features):
        row = feat_table.rows[row_idx + 1]
        bg = "EBF3FB" if row_idx % 2 == 0 else "FFFFFF"
        row.cells[0].text = name
        row.cells[1].text = desc
        for ci in range(2):
            _set_cell_bg(row.cells[ci], bg)
            for para in row.cells[ci].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
        for para in row.cells[0].paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.name = "Courier New"
                run.font.size = Pt(9)

    doc.add_paragraph()
    _add_heading(doc, "2.2. Nhóm 2 – Đặc trưng Host-based (WHOIS)", level=2)
    doc.add_paragraph(
        "domain_age_days: Tuổi của tên miền (ngày từ khi đăng ký). "
        "URL phishing thường dùng domain mới toanh (< 30 ngày). "
        "Gán -1 nếu không tra được.\n"
        "whois_private: 1 nếu thông tin WHOIS bị ẩn (privacy guard), "
        "dấu hiệu mạnh của URL độc hại."
    )

    _add_heading(doc, "2.3. Nhóm 3 – Đặc trưng Content-based (HTTP)", level=2)
    doc.add_paragraph(
        "has_password_form: 1 nếu trang web chứa input type=password "
        "(form đăng nhập giả mạo). Chỉ bật khi fetch_page=True."
    )

    doc.add_page_break()

    # ══════════════════════════════════════════
    # 3. Nhận xét phishing tại Việt Nam
    # ══════════════════════════════════════════
    _add_heading(doc, "3. Nhận xét về phishing tại Việt Nam", level=1)

    vietnam_notes = [
        (
            "3.1. Tên miền và giả mạo thương hiệu",
            (
                "URL phishing Việt Nam ít khi dùng tên miền .vn chính hãng do quy trình đăng ký "
                "tốn kém và cần xác thực. Thay vào đó, kẻ tấn công thường:\n"
                "• Dùng tên miền .com, .net, .top, .xyz rẻ tiền với chuỗi giả tên ngân hàng.\n"
                "• Sử dụng subdomain chứa tên ngân hàng: vietcombank.evil.com, "
                "  techcombank-login.net, bidv-secure.tk, mbbank.phishing.xyz.\n"
                "• Thêm từ 'secure', 'login', 'mobile' vào domain: vietcombanksecure.com."
            ),
        ),
        (
            "3.2. Từ khóa tiếng Việt trong URL",
            (
                "Không ít URL phishing nhắm vào người dùng Việt Nam chứa các chuỗi tiếng Việt "
                "không dấu:\n"
                "• dangnhap, dang-nhap (đăng nhập)\n"
                "• thanhtoan, thanh-toan (thanh toán)\n"
                "• baomat, bao-mat (bảo mật)\n"
                "• taikhoan, tai-khoan (tài khoản)\n"
                "• xacnhan, xac-nhan (xác nhận)\n"
                "• matkhau, mat-khau (mật khẩu)\n"
                "Hệ thống đã tích hợp nhận diện các từ khóa này trong has_sensitive_keyword."
            ),
        ),
        (
            "3.3. Độ dài URL",
            (
                "Nghiên cứu quốc tế cho thấy URL phishing trung bình dài hơn URL benign "
                "(~75 ký tự so với ~30 ký tự). Tại Việt Nam, xu hướng tương tự: "
                "URL giả mạo ngân hàng thường dài 80-150 ký tự do chứa nhiều tham số query "
                "để theo dõi nạn nhân (utm_source, session_id, token, ...)."
            ),
        ),
        (
            "3.4. Các vectơ tấn công phổ biến",
            (
                "• Giả mạo trang đăng nhập Internet Banking (Vietcombank, BIDV, Techcombank).\n"
                "• Lừa đảo qua SMS/Zalo với link rút gọn (bit.ly) dẫn đến trang giả.\n"
                "• Giả mạo trang thương mại điện tử (Shopee, Lazada, Tiki) với tên miền gần giống.\n"
                "• Tấn công qua email phishing giả tên cơ quan nhà nước (BHXH, Công an, ...)."
            ),
        ),
    ]

    for title, content in vietnam_notes:
        _add_heading(doc, title, level=2)
        para = doc.add_paragraph(content)
        para.paragraph_format.space_after = Pt(6)

    doc.add_page_break()

    # ══════════════════════════════════════════
    # 4. Thống kê mô tả đặc trưng
    # ══════════════════════════════════════════
    _add_heading(doc, "4. Thống kê mô tả đặc trưng Lexical", level=1)
    doc.add_paragraph(
        "Bảng dưới đây thể hiện thống kê cơ bản (mean, std, min, max) "
        "cho toàn bộ tập dữ liệu:"
    )

    if stats_df is not None and not stats_df.empty:
        _add_stats_table(doc, stats_df)
    else:
        doc.add_paragraph("(Không có dữ liệu thống kê.)")

    doc.add_paragraph()

    # ══════════════════════════════════════════
    # 5. Kết quả so sánh mô hình
    # ══════════════════════════════════════════
    doc.add_page_break()
    _add_heading(doc, "5. Kết quả so sánh các mô hình ML", level=1)
    doc.add_paragraph(
        "Các mô hình được đánh giá trên tập test (80/20 split, random_state=42) "
        "sử dụng 4 chỉ số: Accuracy, Precision, Recall, F1-Score. "
        "Hàng màu xanh là mô hình tốt nhất theo F1-Score."
    )

    if results_df is not None and not results_df.empty:
        _add_comparison_table(doc, results_df)
    else:
        doc.add_paragraph("(Không có kết quả mô hình.)")

    doc.add_paragraph()

    # ══════════════════════════════════════════
    # 6. Kết luận mô hình tốt nhất
    # ══════════════════════════════════════════
    _add_heading(doc, "6. Kết luận và mô hình được chọn", level=1)

    best_name     = best_model.get("name", "N/A")
    best_feat_set = best_model.get("feature_set", "N/A")
    best_metrics  = best_model.get("metrics", {})

    doc.add_paragraph(
        f"Dựa trên F1-Score tổng hợp, mô hình tốt nhất được chọn là: "
        f"{best_name} với tổ hợp đặc trưng {best_feat_set}."
    )

    conclusion_table = doc.add_table(rows=5, cols=2)
    conclusion_table.style = "Table Grid"
    conclusion_rows = [
        ("Mô hình",       best_name),
        ("Tổ hợp đặc trưng", best_feat_set),
        ("Accuracy",      f"{best_metrics.get('Accuracy',  0):.4f}"),
        ("F1-Score",      f"{best_metrics.get('F1-Score',  0):.4f}"),
        ("Precision",     f"{best_metrics.get('Precision', 0):.4f}"),
    ]
    # Note: table has 5 rows for 5 items above
    for i, (k, v) in enumerate(conclusion_rows):
        row = conclusion_table.rows[i]
        row.cells[0].text = k
        row.cells[1].text = v
        _set_cell_bg(row.cells[0], "DEEAF1")
        for para in row.cells[0].paragraphs:
            for run in para.runs:
                run.bold = True
        bg = "EBF3FB" if i % 2 == 0 else "FFFFFF"
        _set_cell_bg(row.cells[1], bg)

    doc.add_paragraph()
    doc.add_paragraph(
        "Lý do lựa chọn:\n"
        "• F1-Score là chỉ số cân bằng giữa Precision và Recall, đặc biệt quan trọng khi "
        "  dữ liệu mất cân bằng (imbalanced classes).\n"
        "• Trong bài toán phishing detection, Recall cao đảm bảo ít bỏ sót URL độc hại, "
        "  trong khi Precision cao giảm thiểu cảnh báo nhầm gây phiền người dùng.\n"
        "• Mô hình được chọn đạt sự cân bằng tốt nhất giữa hai yếu tố này."
    )

    # ══════════════════════════════════════════
    # Save
    # ══════════════════════════════════════════
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    logger.info(f"[OK] Da luu bao cao: {output_path}")
    return output_path


if __name__ == "__main__":
    # Quick test
    import numpy as np
    results = pd.DataFrame({
        "Model": ["XGBoost", "Decision Tree"],
        "Feature Set": ["X1 (Lexical)", "X1 (Lexical)"],
        "Accuracy":  [0.95, 0.91],
        "Precision": [0.94, 0.90],
        "Recall":    [0.96, 0.92],
        "F1-Score":  [0.95, 0.91],
    })
    best = {
        "name": "XGBoost",
        "feature_set": "X1 (Lexical)",
        "metrics": results.iloc[0].to_dict(),
    }
    stats = pd.DataFrame(np.random.randn(5, 4),
                         index=[f"f{i}" for i in range(5)],
                         columns=["mean", "std", "min", "max"])
    generate_report(
        "test_report.docx", results, best, stats,
        1000, 500, 500, [f"f{i}" for i in range(5)]
    )
